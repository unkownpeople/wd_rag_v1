from __future__ import annotations

"""准备 3 个 DOCX 和 3 个 XLSX 格式 fixture，并生成多格式 manifest。"""

import csv
import hashlib
import json
from pathlib import Path
import shutil


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "annual_reports" / "multiformat_v1" / "raw"
MANIFEST = ROOT / "data" / "annual_reports" / "multiformat_v1" / "manifest.jsonl"
SOURCE_DOCX = ROOT / "data" / "word" / "rag_parse_branch_test.docx"
CSV_SOURCES = [
    ROOT / "data" / "tables_preview" / "finqa_table_01.csv",
    ROOT / "data" / "tables_preview" / "finqa_table_02.csv",
    ROOT / "data" / "tables_preview" / "finqa_table_03.csv",
]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_docx_fixtures() -> list[Path]:
    from docx import Document

    if not SOURCE_DOCX.is_file():
        raise FileNotFoundError(f"DOCX fixture source does not exist: {SOURCE_DOCX}")
    outputs: list[Path] = []
    for index, year in enumerate((2022, 2023, 2024), start=1):
        target = RAW / "docx" / f"word_fixture_{year}.docx"
        target.parent.mkdir(parents=True, exist_ok=True)
        document = Document(SOURCE_DOCX)
        if document.paragraphs:
            document.paragraphs[0].text = f"RAG 多格式解析测试文档 {year}"
        document.add_paragraph(
            f"Fixture document {index}: fiscal year {year}; source_kind=fixture; "
            "paragraph and table locations must remain traceable."
        )
        document.save(target)
        outputs.append(target)
    return outputs


def _write_xlsx_fixtures() -> list[Path]:
    from openpyxl import Workbook

    outputs: list[Path] = []
    for index, source in enumerate(CSV_SOURCES, start=1):
        if not source.is_file():
            raise FileNotFoundError(f"CSV fixture source does not exist: {source}")
        target = RAW / "xlsx" / f"table_fixture_{2021 + index}.xlsx"
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = f"Data_{2021 + index}"
        with source.open("r", encoding="utf-8", newline="") as stream:
            for row in csv.reader(stream):
                sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(target)
        outputs.append(target)
    return outputs


def _manifest_records(docx_files: list[Path], xlsx_files: list[Path]) -> list[dict[str, object]]:
    pdf_manifest = ROOT / "data" / "annual_reports" / "manifest.jsonl"
    pdf_records = [
        json.loads(line)
        for line in pdf_manifest.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    records: list[dict[str, object]] = []
    for item in pdf_records:
        records.append(
            {
                "document_id": item["document_id"],
                "company_id": item["company_id"],
                "company_name": item["company_name"],
                "ticker": item["ticker"],
                "fiscal_year": item["fiscal_year"],
                "report_type": item["report_type"],
                "language": item["language"],
                "currency": item["currency"],
                "unit_scale": item["unit_scale"],
                "source_format": "pdf",
                "source_kind": "original",
                "source_file": item["raw_path"],
                "source_url": item.get("source_url"),
                "source_sha256": item["sha256"],
            }
        )
    for index, path in enumerate(docx_files, start=1):
        year = 2021 + index
        records.append(
            {
                "document_id": f"word_fixture_{year}",
                "company_id": f"fixture_word_{index:02d}",
                "company_name": f"Word Fixture {year}",
                "ticker": None,
                "fiscal_year": year,
                "report_type": "format_fixture",
                "language": "zh-en",
                "currency": None,
                "unit_scale": None,
                "source_format": "docx",
                "source_kind": "fixture",
                "source_file": str(path),
                "source_url": None,
                "source_sha256": _sha256(path),
            }
        )
    for index, path in enumerate(xlsx_files, start=1):
        year = 2021 + index
        records.append(
            {
                "document_id": f"table_fixture_{year}",
                "company_id": f"fixture_xlsx_{index:02d}",
                "company_name": f"XLSX Fixture {year}",
                "ticker": None,
                "fiscal_year": year,
                "report_type": "format_fixture",
                "language": "en",
                "currency": "USD",
                "unit_scale": "as_reported",
                "source_format": "xlsx",
                "source_kind": "fixture",
                "source_file": str(path),
                "source_url": None,
                "source_sha256": _sha256(path),
            }
        )
    return records


def main() -> None:
    docx_files = _write_docx_fixtures()
    xlsx_files = _write_xlsx_fixtures()
    records = _manifest_records(docx_files, xlsx_files)
    records.sort(key=lambda item: (str(item["source_format"]), int(item["fiscal_year"])))
    if len(records) != 9 or {item["source_format"] for item in records} != {"pdf", "docx", "xlsx"}:
        raise RuntimeError("多格式 manifest 未形成 3 PDF + 3 DOCX + 3 XLSX")
    if any(not Path(str(item["source_file"])).is_file() for item in records):
        raise RuntimeError("manifest 存在不存在的 source_file")
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(
        "".join(json.dumps(item, ensure_ascii=False) + "\n" for item in records),
        encoding="utf-8",
    )
    report = {
        "stage": "multiformat_manifest",
        "manifest": str(MANIFEST),
        "document_count": len(records),
        "format_counts": {
            source_format: sum(item["source_format"] == source_format for item in records)
            for source_format in ("pdf", "docx", "xlsx")
        },
        "source_kind_counts": {
            source_kind: sum(item["source_kind"] == source_kind for item in records)
            for source_kind in ("original", "fixture")
        },
        "sha256_missing": sum(not item["source_sha256"] for item in records),
        "passed": len(records) == 9 and all(item["source_sha256"] for item in records),
    }
    (MANIFEST.parent / "manifest.report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
