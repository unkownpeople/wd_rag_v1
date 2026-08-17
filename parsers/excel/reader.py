"""Low-level readers for CSV, TSV, and XLSX rows."""

from __future__ import annotations

import csv
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def cell_value(value: Any) -> Any:
    """Keep typed spreadsheet values in a JSON-compatible form."""

    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def column_name(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def is_empty_row(row: list[Any]) -> bool:
    return not any(value not in (None, "") for value in row)


def read_delimited_rows(input_file: Path, delimiter: str) -> tuple[list[list[Any]], list[str]]:
    warnings: list[str] = []
    try:
        with input_file.open("r", encoding="utf-8-sig", newline="") as file:
            return [list(row) for row in csv.reader(file, delimiter=delimiter)], warnings
    except UnicodeDecodeError:
        warnings.append("UTF-8 decoding failed; retrying with gb18030.")
        with input_file.open("r", encoding="gb18030", newline="") as file:
            return [list(row) for row in csv.reader(file, delimiter=delimiter)], warnings


def read_xlsx_rows(input_file: Path) -> tuple[list[str], list[tuple[str, list[list[Any]]]]]:
    """Read each non-empty worksheet while keeping formulas as formulas."""

    workbook = load_workbook(filename=input_file, read_only=True, data_only=False)
    sheet_names = list(workbook.sheetnames)
    sheet_rows: list[tuple[str, list[list[Any]]]] = []
    for sheet in workbook.worksheets:
        rows = [
            list(row)
            for row in sheet.iter_rows(
                min_row=1,
                max_row=sheet.max_row,
                min_col=1,
                max_col=sheet.max_column,
                values_only=True,
            )
        ]
        while rows and is_empty_row(rows[-1]):
            rows.pop()
        if rows:
            sheet_rows.append((sheet.title, rows))
    workbook.close()
    return sheet_names, sheet_rows
