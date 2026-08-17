from __future__ import annotations

"""Create the local DOCX fixture used by the Word parser branch test."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "word" / "rag_parse_branch_test.docx"


def _set_run_font(run, name: str, size: float, color: str, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            tc_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_width is not None:
                tc_width.set(qn("w:w"), str(widths[index]))
                tc_width.set(qn("w:type"), "dxa")


def _style_paragraph(paragraph, after: float = 6, before: float = 0, line: float = 1.10) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style_paragraph(header, after=0)
    _set_run_font(header.add_run("RAG Parser Fixture | Word branch"), "Calibri", 9, "6B7280")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style_paragraph(footer, after=0)
    _set_run_font(footer.add_run("Local parse test | UTF-8 source"), "Calibri", 9, "6B7280")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _style_paragraph(title, after=4)
    _set_run_font(title.add_run("RAG 多格式解析测试文档"), "Calibri", 23, "000000", bold=True)

    subtitle = document.add_paragraph()
    _style_paragraph(subtitle, after=16)
    _set_run_font(
        subtitle.add_run("Word branch fixture | PDF / Word / Excel parse-only stage"),
        "Calibri",
        12,
        "555555",
    )

    document.add_heading("1. 文本段落", level=1)
    paragraph = document.add_paragraph(
        "这是一篇用于 RAG 解析阶段的本地测试文档。The parser should preserve Chinese text, "
        "English text, section order, punctuation, and numeric facts such as 1024 dimensions "
        "and 19 source pages. 当前步骤只读取原始内容，不执行清洗、切块或向量化。"
    )
    _style_paragraph(paragraph)

    document.add_heading("2. 表格内容", level=1)
    source_note = document.add_paragraph("表 1：保留原始行列关系，用于验证 Word 表格分支。")
    _style_paragraph(source_note, before=0, after=4)
    for run in source_note.runs:
        _set_run_font(run, "Calibri", 10, "6B7280")

    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    _set_table_geometry(table, [2100, 3360, 3900])
    headers = ["字段", "中文值", "English value"]
    rows = [
        ["格式", "Word 文档", "DOCX"],
        ["解析分支", "正文 + 表格", "paragraph + table"],
        ["状态", "待进入清洗", "parse-only"],
    ]
    for column, value in enumerate(headers):
        cell = table.rows[0].cells[column]
        _set_cell_shading(cell, "F2F4F7")
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_paragraph(paragraph, after=0)
        _set_run_font(paragraph.add_run(value), "Calibri", 10, "1F4D78", bold=True)
    for row_index, values in enumerate(rows, start=1):
        for column, value in enumerate(values):
            cell = table.rows[row_index].cells[column]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _style_paragraph(paragraph, after=0)
            _set_run_font(paragraph.add_run(value), "Calibri", 10, "222222")

    document.add_heading("3. 解析预期", level=1)
    expected = document.add_paragraph(
        "Word parser output should contain paragraph records and one table record. "
        "Page numbers remain empty at this stage because DOCX pagination depends on the "
        "renderer; the parser keeps sequence and table number instead."
    )
    _style_paragraph(expected)

    document.core_properties.title = "RAG 多格式解析测试文档"
    document.core_properties.subject = "Word parser branch fixture"
    document.core_properties.author = "RAG project local test"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
